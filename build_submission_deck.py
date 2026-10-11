"""Build the SentinelGate ET AI Hackathon submission PowerPoint."""
from __future__ import annotations
from pathlib import Path
import json

import cv2
from PIL import Image, ImageEnhance
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parent
VIDEO=ROOT/'sentinelgate_vishwajeet_hackathon_submission.mp4'
OUTPUT=ROOT/'Vishwajeet_Kumar_SentinelGate_Pitch_Deck.pptx'
ASSETS=ROOT/'deck_assets';ASSETS.mkdir(exist_ok=True)

BG='071317';SURFACE='0D2229';SURFACE2='091A20';LINE='264149';TEXT='EDF5F1';MUTED='8CA39F';GREEN='69E3AC';AMBER='F2BD63';RED='FF747C';BLUE='6FB9FF';WHITE='FFFFFF'
SW,SH=13.333,7.5


def rgb(value):return RGBColor.from_string(value)


def extract_frame(second,name,darken=1.0):
    path=ASSETS/name
    capture=cv2.VideoCapture(str(VIDEO));capture.set(cv2.CAP_PROP_POS_MSEC,second*1000);ok,frame=capture.read();capture.release()
    if not ok:raise RuntimeError(f'Could not extract video frame at {second}s')
    cv2.imwrite(str(path),frame)
    if darken!=1.0:
        image=Image.open(path).convert('RGB');image=ImageEnhance.Brightness(image).enhance(darken);image.save(path)
    return path


def add_bg(slide,color=BG):
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(SW),Inches(SH));shape.fill.solid();shape.fill.fore_color.rgb=rgb(color);shape.line.fill.background()


def add_text(slide,value,x,y,w,h,size=18,color=TEXT,bold=False,font='Aptos',align=PP_ALIGN.LEFT,margin=0,anchor=MSO_ANCHOR.TOP):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=box.text_frame;frame.clear();frame.word_wrap=True;frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=Inches(margin);frame.vertical_anchor=anchor
    paragraph=frame.paragraphs[0];paragraph.text=value;paragraph.alignment=align;paragraph.font.name=font;paragraph.font.size=Pt(size);paragraph.font.bold=bold;paragraph.font.color.rgb=rgb(color)
    return box


def add_rich(slide,runs,x,y,w,h,size=18,align=PP_ALIGN.LEFT):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=box.text_frame;frame.clear();frame.word_wrap=True
    paragraph=frame.paragraphs[0];paragraph.alignment=align
    for value,color,bold in runs:
        run=paragraph.add_run();run.text=value;run.font.name='Aptos';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=rgb(color)
    return box


def box(slide,x,y,w,h,fill=SURFACE,line=LINE,radius=True,width=1):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill);shape.line.color.rgb=rgb(line);shape.line.width=Pt(width)
    return shape


def line(slide,x1,y1,x2,y2,color=LINE,width=1):
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x1),Inches(y1),Inches(max(.01,x2-x1)),Inches(max(.01,y2-y1)));shape.fill.solid();shape.fill.fore_color.rgb=rgb(color);shape.line.fill.background();return shape


def header(slide,section,title,number):
    add_text(slide,'AGENTSHIELD',.45,.25,2.4,.25,12,GREEN,True,'Cascadia Mono')
    add_text(slide,'ET AI HACKATHON · AGENTIC EDITION',10.15,.26,2.72,.2,8,MUTED,False,'Cascadia Mono',PP_ALIGN.RIGHT)
    add_text(slide,section.upper(),.55,.78,5.5,.25,10,GREEN,True,'Cascadia Mono')
    add_text(slide,title,.55,1.05,12.1,.56,27,TEXT,True,'Aptos Display')
    line(slide,.55,1.68,12.78,1.69,LINE,1)
    add_text(slide,f'{number:02}',12.43,7.13,.35,.18,9,MUTED,True,'Cascadia Mono',PP_ALIGN.RIGHT)


def tag(slide,label,x,y,color=GREEN,w=None):
    if w is None:w=max(1.1,.09*len(label)+.42)
    shape=box(slide,x,y,w,.34,SURFACE2,color,True,1);add_text(slide,label.upper(),x,y+.08,w,.15,8,color,True,'Cascadia Mono',PP_ALIGN.CENTER);return shape


def metric(slide,value,label,x,y,w=1.85,color=GREEN):
    box(slide,x,y,w,1.0,SURFACE,LINE);add_text(slide,label.upper(),x+.16,y+.16,w-.32,.18,8,MUTED,True,'Cascadia Mono');add_text(slide,value,x+.16,y+.48,w-.32,.33,22,color,True,'Cascadia Mono')


def bullets(slide,items,x,y,w,h,size=16,color=TEXT,accent=GREEN):
    container=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=container.text_frame;frame.clear();frame.word_wrap=True;frame.margin_left=Inches(.02)
    for index,item in enumerate(items):
        p=frame.paragraphs[0] if index==0 else frame.add_paragraph();p.text=item;p.font.name='Aptos';p.font.size=Pt(size);p.font.color.rgb=rgb(color);p.space_after=Pt(10);p.level=0;p.text='  '+p.text
    for index,p in enumerate(frame.paragraphs):
        marker=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y+.08+index*.47),Inches(.11),Inches(.11));marker.fill.solid();marker.fill.fore_color.rgb=rgb(accent);marker.line.fill.background()
    return container


def picture(slide,path,x,y,w,h,border=LINE):
    box(slide,x-.02,y-.02,w+.04,h+.04,SURFACE2,border,True,1);slide.shapes.add_picture(str(path),Inches(x),Inches(y),Inches(w),Inches(h))


def pipeline(slide,items,y=3.35):
    count=len(items);gap=.19;total=12.05;w=(total-gap*(count-1))/count
    for i,(title,subtitle,color) in enumerate(items):
        x=.65+i*(w+gap);box(slide,x,y,w,1.15,SURFACE2,color,True,1.5);add_text(slide,f'{i+1:02}',x+.13,y+.13,.3,.2,9,color,True,'Cascadia Mono');add_text(slide,title,x+.13,y+.42,w-.26,.24,13,TEXT,True);add_text(slide,subtitle,x+.13,y+.75,w-.26,.22,9,MUTED,False,'Cascadia Mono')
        if i<count-1:add_text(slide,'→',x+w,y+.43,gap,.25,15,GREEN,True,'Cascadia Mono',PP_ALIGN.CENTER)


def build():
    frames={
        'title':extract_frame(12,'01_title.png',.72),'risk':extract_frame(47,'02_risk.png'),'firewall':extract_frame(82,'03_firewall.png'),
        'lab':extract_frame(112,'04_lab.png'),'tracker':extract_frame(166,'05_tracker.png'),'neutralize':extract_frame(207,'06_neutralize.png'),
        'handoff':extract_frame(232,'07_handoff.png'),'agentic':extract_frame(263,'08_agentic.png'),'egress':extract_frame(293,'09_egress.png'),
        'evaluation':extract_frame(327,'10_evaluation.png'),'architecture':extract_frame(360,'11_architecture.png')}
    benchmark=json.loads((ROOT/'benchmark_results.json').read_text(encoding='utf-8'));events=json.loads((ROOT/'file_tracking_events.json').read_text(encoding='utf-8'))
    prs=Presentation();prs.slide_width=Inches(SW);prs.slide_height=Inches(SH);blank=prs.slide_layouts[6]
    props=prs.core_properties
    props.author='Vishwajeet Kumar'
    props.last_modified_by='Vishwajeet Kumar'
    props.title='SentinelGate - Zero-Trust Prompt Injection Firewall'
    props.subject='ET AI Hackathon: Agentic Edition | Presented by Accenture'
    props.keywords='Agentic AI, cybersecurity, prompt injection, zero trust, SentinelGate, AgentShield'
    props.comments='Submission pitch deck for Problem Statement 2: Agentic Cybersecurity - Prompt Injection Firewall.'

    # 1 — title
    slide=prs.slides.add_slide(blank);add_bg(slide);slide.shapes.add_picture(str(frames['title']),0,0,Inches(SW),Inches(SH))
    shade=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(7.4),Inches(SH));shade.fill.solid();shade.fill.fore_color.rgb=rgb(BG);shade.line.fill.background()
    add_text(slide,'ET AI HACKATHON: AGENTIC EDITION',.65,.55,6.2,.3,12,GREEN,True,'Cascadia Mono')
    add_text(slide,'Presented by Accenture',.65,.92,4,.25,12,MUTED)
    add_text(slide,'SentinelGate',.65,1.58,5.9,.65,38,TEXT,True,'Aptos Display')
    add_text(slide,'Zero-Trust Prompt Injection Firewall',.65,2.27,6.1,.5,23,GREEN,True)
    add_text(slide,'Inspect before the agent can act.',.65,3.02,5.4,.38,19,TEXT)
    tag(slide,'Problem Statement 2',.65,3.73,GREEN,1.75);tag(slide,'F3 · 9/9 attacks',2.58,3.73,BLUE,1.65);tag(slide,'D2 verified',4.41,3.73,AMBER,1.35)
    add_text(slide,'VISHWAJEET KUMAR',.65,5.7,5.4,.35,19,TEXT,True,'Cascadia Mono')
    add_text(slide,'Working prototype + structural architecture',.65,6.16,5.5,.3,14,MUTED)

    # 2 — problem
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'The challenge','Untrusted content can become agent authority',2)
    add_text(slide,'THE BRIEF',.65,1.98,2,.2,10,RED,True,'Cascadia Mono')
    add_text(slide,'Detect and neutralize malicious instructions before they influence an AI agent—while legitimate content passes with minimal disruption.',.65,2.3,5.35,1.25,22,TEXT,True)
    add_text(slide,'WHY IT MATTERS',.65,4.08,2,.2,10,AMBER,True,'Cascadia Mono')
    bullets(slide,['Secret and credential extraction','Unauthorized tool execution','Context poisoning across turns','Indirect attacks hidden in retrieved content'],.65,4.42,5.25,1.8,15,TEXT,RED)
    picture(slide,frames['risk'],6.45,2.02,6.25,3.52,RED)
    add_text(slide,'INPUT SURFACE',6.45,5.84,2,.18,9,MUTED,True,'Cascadia Mono')
    add_text(slide,'Chat · Web · Email · PDF · DOCX · HTML · API · Code · OCR',6.45,6.15,6.15,.3,14,TEXT,True)

    # 3 — solution
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'The solution','Security decisions live outside the model',3)
    add_text(slide,'SentinelGate is a security control plane between untrusted content and trusted agent execution.',.65,2.0,11.9,.5,22,TEXT,True)
    pipeline(slide,[('Extract','multiformat + hidden content',BLUE),('Decode','Unicode · Base64 · hex',AMBER),('Detect','rules + semantic + state',GREEN),('Decide','allow · quarantine · block',RED),('Enforce','signed context + tool gate',GREEN)],3.05)
    differentiators=[('Signed context broker','Blocked content becomes an empty payload—not a warning the agent may ignore.'),('Immutable intent anchor','Detects semantic drift and multi-turn jailbreaks across the session.'),('Least-privilege tool gate','Server-side permissions, strict schemas, destination policy, confirmation.'),('Output + evidence guard','Redaction, honey-canaries, and a verifiable SHA-256 audit chain.')]
    for i,(title,body) in enumerate(differentiators):
        x=.65+(i%2)*6.1;y=4.82+(i//2)*.91;box(slide,x,y,5.82,.72,SURFACE,LINE);add_text(slide,title,x+.18,y+.13,2.35,.2,12,GREEN,True);add_text(slide,body,x+2.43,y+.12,3.17,.42,11,MUTED)

    # 4 — coverage
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'Requirement coverage','F3 implemented · D2 verified · D3 conditional',4)
    add_rich(slide,[('F3 ',GREEN,True),('requires ≥7 attack families. SentinelGate demonstrates ',TEXT,False),('9/9.',GREEN,True)],.65,1.95,7.1,.35,17)
    names=['Instruction Override','Role Change','Secret Extraction','Tool Abuse','Credential Theft','Context Poisoning','Multi-Step Jailbreak','Encoded Instructions','Indirect Injection']
    for i,name in enumerate(names):
        x=.65+(i%3)*2.42;y=2.53+(i//3)*.75;box(slide,x,y,2.2,.58,SURFACE2,GREEN);add_text(slide,f'{i+1:02}',x+.12,y+.18,.28,.15,8,GREEN,True,'Cascadia Mono');add_text(slide,name,x+.45,y+.16,1.62,.23,10,TEXT,True)
    picture(slide,frames['lab'],8.12,2.05,4.58,2.58,GREEN)
    box(slide,8.12,4.92,4.58,1.36,SURFACE,AMBER);add_text(slide,'DEPTH CLAIM',8.35,5.13,1.4,.18,9,AMBER,True,'Cascadia Mono');add_text(slide,'D2 verified locally',8.35,5.47,2.0,.28,18,TEXT,True);add_text(slide,'D3 remains conditional until image OCR is runtime-ready and demonstrated.',8.35,5.85,3.95,.32,11,MUTED)

    # 5 — architecture
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'Structural architecture','Trust boundaries and enforcement flow',5)
    add_text(slide,'UNTRUSTED INPUT ZONE',.65,1.98,2.5,.2,9,RED,True,'Cascadia Mono');add_text(slide,'TRUSTED EXECUTION ZONE',10.03,1.98,2.7,.2,9,GREEN,True,'Cascadia Mono',PP_ALIGN.RIGHT)
    pipeline(slide,[('Adapters','chat · files · OCR',RED),('Normalizer','canonical text',BLUE),('Ensemble','rules · semantic · state',GREEN),('Orchestrator','risk + policy',AMBER),('Context broker','signed data envelope',GREEN),('Agent + tools','least privilege',BLUE)],2.55)
    line(slide,.75,4.18,12.55,4.2,LINE)
    controls=[('01','Intent contract','Immutable session objective'),('02','Tool gate','Authorization + schemas'),('03','Output guard','Secrets + PII + canary'),('04','Evidence ledger','Hash-chained decisions')]
    for i,(num,title,body) in enumerate(controls):
        x=.65+i*3.05;box(slide,x,4.55,2.78,1.18,SURFACE,LINE);add_text(slide,num,x+.16,4.74,.3,.2,9,GREEN,True,'Cascadia Mono');add_text(slide,title,x+.55,4.71,1.9,.25,13,TEXT,True);add_text(slide,body,x+.16,5.18,2.42,.25,10,MUTED)
    add_text(slide,'Fail-closed rule: residual malicious content or policy failure never reaches the agent.',.65,6.22,11.9,.35,15,GREEN,True,align=PP_ALIGN.CENTER)

    # 6 — live firewall
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'Working prototype','Live firewall: mixed content → enforceable verdict',6)
    picture(slide,frames['firewall'],.65,2.0,7.2,4.05,RED)
    add_text(slide,'WHAT THE DEMO PROVES',8.2,2.08,3.8,.2,10,GREEN,True,'Cascadia Mono')
    bullets(slide,['Useful revenue data is preserved','Instruction override is detected','Secret extraction is detected','Credential theft is detected','Provenance escalates indirect risk','Decision includes evidence and route'],8.2,2.48,4.3,2.9,15)
    tag(slide,'BLOCK',8.2,5.6,RED,1.0);tag(slide,'100 / 100 risk',9.38,5.6,AMBER,1.5);tag(slide,'post-scan verified',11.06,5.6,GREEN,1.5)

    # 7 — tracker
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'Live enterprise defense','Real-time malicious file tracking and quarantine',7)
    picture(slide,frames['tracker'],.55,1.92,8.35,4.7,BLUE)
    allowed=sum(e['verdict']=='ALLOW' for e in events);blocked=sum(e['verdict']=='BLOCK' for e in events)
    metric(slide,str(len(events)),'files inspected',9.25,2.0,1.52,BLUE);metric(slide,str(allowed),'allowed',10.95,2.0,1.52,GREEN)
    metric(slide,str(blocked),'quarantined',9.25,3.2,1.52,RED);metric(slide,'SHA-256','fingerprinted',10.95,3.2,1.52,AMBER)
    add_text(slide,'REAL PIPELINE EVENTS',9.25,4.62,2.8,.2,9,GREEN,True,'Cascadia Mono')
    bullets(slide,['Hidden HTML instruction','Encoded API injection','Credential-theft email','Clean report + code pass'],9.25,4.98,3.15,1.35,12,TEXT,RED)

    # 8 — agentic security
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'Agentic capability','Stateful intent defense and action governance',8)
    picture(slide,frames['agentic'],.65,1.98,6.55,3.68,GREEN)
    cards=[('INTENT ANCHOR','Locks the declared objective and detects drift across turns.',GREEN),('TOOL POLICY','Validates authorization, allowlist, schema, destination, and payload.',BLUE),('HUMAN OVERSIGHT','High-risk actions stay disabled until explicit confirmation.',AMBER)]
    for i,(title,body,color) in enumerate(cards):
        y=2.02+i*1.22;box(slide,7.55,y,5.0,.95,SURFACE,color);add_text(slide,title,7.78,y+.16,1.72,.18,9,color,True,'Cascadia Mono');add_text(slide,body,9.3,y+.13,2.98,.48,12,TEXT)
    add_text(slide,'The model proposes. The security plane authorizes.',.65,6.08,11.9,.35,18,GREEN,True,align=PP_ALIGN.CENTER)

    # 9 — handoff + egress
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'End-to-end enforcement','No unsafe handoff. No unsafe output.',9)
    picture(slide,frames['handoff'],.65,2.0,5.9,3.32,GREEN);picture(slide,frames['egress'],6.78,2.0,5.9,3.32,BLUE)
    add_text(slide,'INBOUND CONTROL',.65,5.58,2,.18,9,GREEN,True,'Cascadia Mono');add_text(slide,'Blocked payloads contain no content and no capabilities.',.65,5.91,5.7,.42,13,TEXT,True)
    add_text(slide,'EGRESS CONTROL',6.78,5.58,2,.18,9,BLUE,True,'Cascadia Mono');add_text(slide,'Secrets, PII, system prompts, and honey-canaries are redacted.',6.78,5.91,5.7,.42,13,TEXT,True)

    # 10 — evaluation
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'Measured evidence','Transparent evaluation—with limitations stated',10)
    values=[(f"{benchmark['samples']:,}",'cases',BLUE),(f"{benchmark['recall']*100:.0f}%",'synthetic recall',GREEN),(f"{benchmark['precision']*100:.0f}%",'synthetic precision',GREEN),(f"{benchmark['false_positive_rate']*100:.0f}%",'false positives',GREEN),(f"{benchmark['p95_ms']:.1f} ms",'P95 latency',AMBER),('13 / 13','integration tests',GREEN)]
    for i,(value,label,color) in enumerate(values):metric(slide,value,label,.65+(i%3)*2.32,2.04+(i//3)*1.27,2.05,color)
    picture(slide,frames['evaluation'],7.85,2.03,4.82,2.71,GREEN)
    box(slide,7.85,5.04,4.82,1.18,SURFACE,AMBER);add_text(slide,'HONEST LIMITATION',8.08,5.23,1.8,.18,9,AMBER,True,'Cascadia Mono');add_text(slide,'Synthetic regression evidence—not production accuracy. Independent benchmark validation is the next gate.',8.08,5.57,4.24,.46,12,TEXT,True)
    add_text(slide,'Coverage: 45 attack phrasings · 10 context wrappers · PDF/DOCX/EML/HTML adapter tests',.65,5.92,6.85,.45,12,MUTED)

    # 11 — impact
    slide=prs.slides.add_slide(blank);add_bg(slide);header(slide,'Impact and scalability','From hackathon prototype to enterprise security plane',11)
    impacts=[('RISK REDUCTION','Stops malicious influence before context or tool execution.',RED),('PRODUCTIVITY','Preserves legitimate content instead of blocking entire workflows.',GREEN),('AUDITABILITY','Creates evidence for security review, compliance, and incident response.',BLUE)]
    for i,(title,body,color) in enumerate(impacts):
        x=.65+i*4.05;box(slide,x,2.02,3.75,1.2,SURFACE,color);add_text(slide,title,x+.2,2.23,2.4,.2,10,color,True,'Cascadia Mono');add_text(slide,body,x+.2,2.61,3.25,.42,13,TEXT)
    add_text(slide,'PRODUCTION ROADMAP',.65,3.78,2.5,.2,10,GREEN,True,'Cascadia Mono')
    roadmap=[('1','Independent benchmark','Calibrated thresholds + confidence intervals'),('2','Model upgrade','ONNX embeddings + adversarial retraining'),('3','Enterprise controls','RBAC · tenant isolation · SIEM export'),('4','D3 completion','OCR runtime + representative image tests')]
    for i,(num,title,body) in enumerate(roadmap):
        x=.65+i*3.05;box(slide,x,4.18,2.77,1.4,SURFACE2,LINE);add_text(slide,num,x+.15,4.37,.25,.2,10,GREEN,True,'Cascadia Mono');add_text(slide,title,x+.52,4.34,2.0,.25,13,TEXT,True);add_text(slide,body,x+.15,4.86,2.4,.42,11,MUTED)
    add_text(slide,'Designed as an adapter-based, policy-versioned control plane that scales independently of the agent.',.65,6.12,11.9,.35,14,GREEN,True,align=PP_ALIGN.CENTER)

    # 12 — close
    slide=prs.slides.add_slide(blank);add_bg(slide)
    add_text(slide,'VISHWAJEET KUMAR',.65,.52,4.0,.25,12,GREEN,True,'Cascadia Mono')
    add_text(slide,'SentinelGate',.65,1.3,6.1,.65,40,TEXT,True,'Aptos Display');add_text(slide,'Security outside the model.',.65,2.03,6.4,.5,25,GREEN,True)
    add_text(slide,'Untrusted content becomes controlled data—not agent authority.',.65,2.74,6.3,.6,19,MUTED)
    reasons=[('Significant','Protects real agent actions'),('Innovative','Signed context + intent anchors'),('Agentic','State, tools, recovery, oversight'),('Robust','Fail closed + tamper-evident evidence')]
    for i,(title,body) in enumerate(reasons):
        x=.65+(i%2)*3.18;y=3.63+(i//2)*.95;box(slide,x,y,2.9,.72,SURFACE,LINE);add_text(slide,title,x+.16,y+.13,1.1,.18,10,GREEN,True,'Cascadia Mono');add_text(slide,body,x+1.15,y+.11,1.55,.35,11,TEXT)
    box(slide,7.5,.75,5.15,5.75,SURFACE2,GREEN);add_text(slide,'SUBMISSION CHECKLIST',7.83,1.08,3.8,.25,12,GREEN,True,'Cascadia Mono')
    add_text(slide,'✓  Working prototype',7.83,1.68,4.0,.3,17,TEXT,True);add_text(slide,'✓  Pitch deck',7.83,2.2,4.0,.3,17,TEXT,True);add_text(slide,'✓  Narrated demo video',7.83,2.72,4.0,.3,17,TEXT,True)
    box(slide,7.83,3.36,4.46,.98,SURFACE,GREEN)
    add_text(slide,'GITHUB REPOSITORY',8.05,3.51,2.2,.18,9,GREEN,True,'Cascadia Mono')
    repo_link=add_text(slide,'github.com/vishwajeetkumar9631/\nsentinelgate-agentic-security',8.05,3.78,3.95,.48,12,TEXT,True)
    repo_link.click_action.hyperlink.address='https://github.com/vishwajeetkumar9631/sentinelgate-agentic-security'
    add_text(slide,'Final claim',7.83,4.72,1.4,.18,9,MUTED,True,'Cascadia Mono');add_text(slide,'F3 implemented · D2 verified · D3 conditional',7.83,5.05,4.2,.5,17,GREEN,True)
    add_text(slide,'Inspect before the agent can act.',7.83,5.85,4.2,.3,16,TEXT,True)
    add_text(slide,'ET AI Hackathon: Agentic Edition · Presented by Accenture',.65,6.82,8.0,.22,10,MUTED)

    prs.save(OUTPUT)
    return OUTPUT


if __name__=='__main__':print(build())
